"""盲标表：生成带下拉框的 xlsx，并把填好的结果读回 labels.json。

表里只给公司、岗位、城市、JD 要点和全文，不给任何 AI 结论或主表状态，避免锚定。
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..jd import clean_jd, split_clauses
from .pool import PoolEntry

JUDGMENTS = ["投", "不投", "说不准"]
REASONS = ["能力匹配", "硬门槛不符", "核心技能缺口", "方向不想做", "城市或其他偏好", "其他"]
LABEL_SHEET = "标注"
FULL_SHEET = "JD全文"

_DUTY = re.compile(r"职责|岗位介绍|职位描述|工作内容|岗位描述|岗位说明|Duty|Responsibilit", re.I)
_REQ = re.compile(r"要求|任职资格|招聘条件|岗位条件|任职条件|应聘条件|Require|Qualification", re.I)
_BONUS = re.compile(r"加分项|优先条件|加分")
_SKIP = re.compile(r"福利|薪酬|待遇|公司信息|公司介绍|公司简介|工作地点|职位类别|团队介绍|关于我们")
_NUMBERING = re.compile(
    r"^\s*(?:[（(]?\d{1,2}(?:-\d{1,2})?[)）.、．]|[（(]?[一二三四五六七八九十]+[)）、.．]|[-*•·●○▪◆■★])\s*"
)
_PAREN = re.compile(r"[（(][^（）()]*[）)]")
_KEYED = re.compile(r"^(.{1,16}?)[：:]\s*(.+)$")
_SECTION_KEY = re.compile(r"^(?:岗位|职位|任职|工作|招聘|应聘)")


def _header_kind(line: str) -> str | None:
    """节标题的类别；「岗位职责 (研究方向择一)：」「【职责 description】」「2、技能要求」都算。"""
    core = _PAREN.sub("", _NUMBERING.sub("", line)).strip().strip("【】[]#* ").rstrip("：: ")
    if len(core) > 16 or "：" in core or ":" in core:  # 「学历要求：本科」是要求条目，不是节标题
        return None
    for kind, pattern in (("bonus", _BONUS), ("duty", _DUTY), ("req", _REQ), ("skip", _SKIP)):
        if pattern.search(core):
            return kind
    return None


def _clip(line: str, limit: int = 60) -> str:
    line = _NUMBERING.sub("", line).strip().rstrip("；;。")
    return line if len(line) <= limit else line[: limit - 1] + "…"


def _split_pick(lines: list[str], limit: int) -> tuple[list[str], list[str]]:
    # 同一节里有编号条目时只取编号条目，跳过「团队介绍」之类的引言和「学历」「技能」这类小标题
    items = [line for line in lines if _NUMBERING.match(line)]
    pool = items if len(items) >= 2 else lines
    return pool[:limit], pool[limit:]


def jd_points(text: str, duties: int = 2, requirements: int = 3) -> list[str]:
    """启发式抽取 JD 要点：职责前 2 条 + 要求前 3 条；分节识别不全时用其余正文补到 5 条。"""
    sections: dict[str, list[str]] = {"duty": [], "req": [], "bonus": [], "skip": [], "other": []}
    current = "other"
    for line in split_clauses(text):
        kind = _header_kind(line)
        if kind:
            current = kind
            continue
        keyed = _KEYED.match(line)
        if keyed and _SECTION_KEY.match(keyed.group(1)) and (kind := _header_kind(keyed.group(1))):
            current, line = kind, keyed.group(2)  # 「任职要求：本科及以上……」标题和正文在同一行
        sections[current].append(line)
    duty_pick, duty_rest = _split_pick(sections["duty"], duties)
    req_pick, req_rest = _split_pick(sections["req"], requirements)
    picked = duty_pick + req_pick
    for extra in (duty_rest, req_rest, sections["other"]):
        picked += extra[: max(0, duties + requirements - len(picked))]
    return [point for point in (_clip(line) for line in picked) if point]


def display_city(city: str) -> str:
    city = (city or "").strip()
    if not city or city.startswith("官方未"):
        return "未标注"
    return city if len(city) <= 24 else city[:23] + "…"


def write_sheet(entries: list[PoolEntry], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = LABEL_SHEET
    headers = ["序号", "公司", "岗位", "城市", "JD 要点（点序号看全文）", "判断", "理由类别", "备注（一句话即可）", "编号"]
    widths = [6, 18, 26, 12, 72, 9, 15, 30, 13]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for col, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "B2"

    full = wb.create_sheet(FULL_SHEET)
    full.append(["序号", "公司", "岗位", "JD 全文"])
    for col, width in enumerate([6, 18, 26, 110], start=1):
        full.cell(row=1, column=col).font = Font(bold=True)
        full.column_dimensions[full.cell(row=1, column=col).column_letter].width = width
    full.freeze_panes = "B2"

    top_wrap = Alignment(vertical="top", wrap_text=True)
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for index, entry in enumerate(entries, start=1):
        points = jd_points(entry.jd_text)
        row = index + 1
        ws.append([index, entry.company, entry.title, display_city(entry.city), "\n".join(f"• {p}" for p in points), None, None, None, entry.eval_id])
        link = ws.cell(row=row, column=1)
        link.hyperlink = f"#'{FULL_SHEET}'!A{row}"
        link.font = Font(color="0563C1", underline="single")
        for col in range(1, 10):
            ws.cell(row=row, column=col).alignment = top_wrap
        for col in (6, 7, 8):
            ws.cell(row=row, column=col).fill = input_fill
        ws.row_dimensions[row].height = max(15 * max(len(points), 2), 30)

        full_text = clean_jd(entry.jd_text)
        full.append([index, entry.company, entry.title, full_text])
        for col in range(1, 5):
            full.cell(row=row, column=col).alignment = top_wrap
        lines = sum(max(1, len(line) // 60 + 1) for line in full_text.splitlines())
        full.row_dimensions[row].height = min(15 * lines, 409)

    last = len(entries) + 1
    judgment = DataValidation(type="list", formula1='"' + ",".join(JUDGMENTS) + '"', allow_blank=True)
    reason = DataValidation(type="list", formula1='"' + ",".join(REASONS) + '"', allow_blank=True)
    judgment.error = reason.error = "请从下拉框里选"
    ws.add_data_validation(judgment)
    ws.add_data_validation(reason)
    judgment.add(f"F2:F{last}")
    reason.add(f"G2:G{last}")
    ws.column_dimensions["I"].hidden = True

    guide = wb.create_sheet("填写说明", 0)
    lines = [
        "这张表记录你的投递意向，供离线对比；它不是能力匹配或录用率的标准答案。",
        "",
        "每行只填黄色三格：",
        "1. 先看「JD全文」页，再判断是否值得尝试，选「投 / 不投 / 说不准」。",
        "   投＝愿意做、没有已知硬条件冲突、核心工作有直接或可迁移经历，愿意承担补短板的成本。",
        "   不投＝明确硬条件冲突，或本人不接受方向/地域，或核心工作不打算做或补。",
        "   说不准＝关键条件或本人经历信息不足，暂时不能判断；不因一个陌生关键词就判不投。",
        "2. 理由类别：选最主要的一个。",
        "   能力匹配＝经历对得上；硬门槛不符＝学历/专业/届别/证书等不满足；核心技能缺口＝关键技术没做过；",
        "   方向不想做＝能做但不想走这个方向；城市或其他偏好＝城市、实习、公司规模等个人偏好；其他＝写在备注里。",
        "3. 备注：一句话，可空。例如「要 CUDA，没做过」「杭州，方向对口」。",
        "",
        "注意：",
        "• 按你自己的判断填，不用管这些岗位以前是否登记过、AI 以前怎么评。表里刻意没有放 AI 结论。",
        "• JD 要点含省略号，且只选取部分条目，不能代替全文。点「序号」跳到「JD全文」页核对硬条件。",
        "• 行顺序已打乱。不要删行、不要改最后一列隐藏的编号；可以排序。",
        "• 填完保存即可，文件名不用改。约 30–40 分钟。",
    ]
    for line in lines:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True)
    wb.active = 1  # 打开时停在「标注」页
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_labels(path: Path) -> dict[str, dict]:
    """读回填好的表：{eval_id: {judgment, reason, note}}，未填判断的行跳过。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    ws = workbook[LABEL_SHEET]
    labels: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[8]:
            continue
        judgment = (row[5] or "").strip() if isinstance(row[5], str) else row[5]
        if judgment in (None, ""):
            continue
        if judgment not in JUDGMENTS:
            workbook.close()
            raise ValueError("标注中存在无效判断，请使用投/不投/说不准")
        if str(row[8]) in labels:
            workbook.close()
            raise ValueError("标注表存在重复 eval_id")
        labels[str(row[8])] = {
            "judgment": judgment,
            "reason": (row[6] or "").strip() if isinstance(row[6], str) else "",
            "note": (row[7] or "").strip() if isinstance(row[7], str) else "",
        }
    workbook.close()
    return labels


def write_labels(labels: dict[str, dict], path: Path) -> None:
    path.write_text(json.dumps(labels, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
